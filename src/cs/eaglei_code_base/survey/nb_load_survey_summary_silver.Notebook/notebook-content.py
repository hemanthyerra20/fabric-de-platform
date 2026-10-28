# Fabric notebook source

# METADATA ********************

# META {
# META   "kernel_info": {
# META     "name": "synapse_pyspark"
# META   },
# META   "dependencies": {
# META     "lakehouse": {
# META       "default_lakehouse": "0277ffad-3784-4551-8d7d-66d1b8ff633a",
# META       "default_lakehouse_name": "lh_cs_bronze",
# META       "default_lakehouse_workspace_id": "f2aea860-156d-4e18-b078-1356be7161de",
# META       "known_lakehouses": [
# META         {
# META           "id": "0277ffad-3784-4551-8d7d-66d1b8ff633a"
# META         },
# META         {
# META           "id": "08846ef1-977d-4698-b06f-84fdbd834bc6"
# META         }
# META       ]
# META     },
# META     "environment": {
# META       "environmentId": "f45c2c8e-a6e5-4098-aaf6-567ab6b6ab98",
# META       "workspaceId": "00000000-0000-0000-0000-000000000000"
# META     }
# META   }
# META }

# CELL ********************

%run nb_pipeline_monitoring_utils

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

%run nb_utils

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

VAR_CS_LIB_NAME = "var_cs_process"
MASTER_RUN_ID = ""
RUN_ID = ""
PIPELINE_NAME = ""
NOTEBOOK_NAME = ""
DOMAIN = ""
SOURCE_NAME = ""
ACTIVITY_NAME = ""
LEVEL = ""
LAYER = ""
ENVIRONMENT = ""
WORKSPACE_NAME = ""
WORKSPACE_ID = ""

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

workspace_id = get_variable(VAR_CS_LIB_NAME, "cs_store_workspace_id")
silver_lakehouse_id = get_variable(VAR_CS_LIB_NAME, "lh_cs_silver_id")
gold_lakehouse_id = get_variable(VAR_CS_LIB_NAME, "lh_cs_gold_id")
silver_schema_name = "survey"
start_time = datetime.now()
source_table = "silver_tables"
target_table = f"lh_cs_silver.{silver_schema_name}.graduate_outcomes_summary"

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

import pandas as pd
import numpy as np
from openpyxl import load_workbook
from pyspark.sql import SparkSession
from pyspark.sql.types import *
 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

SOURCE_FILE = (
    "/lakehouse/default/Files/survey/"
    "Graduate Outcomes Summary Data_0509.xlsx"
)
OUTPUT_EXCEL = (
    "/lakehouse/default/Files/survey/"
    "NWTC_Data_test.xlsx"
)


# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

wb = load_workbook(
    filename=SOURCE_FILE,
    data_only=True
)
sheet_names = wb.sheetnames
print(sheet_names)
 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

nwtc_sheets = [
    s
    for s in sheet_names
    if ("NWTC Data" in s)
    and ("Program Data" not in s)
    and (s != "NWTC Data Ref")
]
print("Sheets being processed:")
for s in nwtc_sheets:
    print(s)
 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

SUBJECTS = [
    "NWTC Graduate Follow-Up Data",
    "Employment",
    "Location",
    "Satisfaction",
    "Why grads choose NWTC"
]
final_rows = []
for sheet_name in nwtc_sheets:
    year = sheet_name.replace(
        " NWTC Data",
        ""
    )
    ws = wb[sheet_name]
    values = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                values.append(cell)
    current_subject = None
    i = 0
    while i < len(values):
        current_value = values[i]
        if str(current_value) in SUBJECTS:
            current_subject = str(current_value)
            i += 1
            continue
        if str(current_value) == f"{year} Grads":
            i += 1
            continue
        if (
            current_subject is not None
            and i + 1 < len(values)
        ):
            measure_name = str(values[i])
            measure_value = values[i + 1]
            final_rows.append(
                (
                    year,
                    current_subject,
                    measure_name,
                    measure_value
                )
            )
            i += 2
        else:
            i += 1

 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

schema = StructType([
    StructField(
        "academic_year",
        StringType(),
        True
    ),
    StructField(
        "subject_area",
        StringType(),
        True
    ),
    StructField(
        "measure_name",
        StringType(),
        True
    ),
    StructField(
        "measure_value",
        StringType(),
        True
    )
])
df = spark.createDataFrame(
    final_rows,
    schema
)
display(df)
 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

df.groupBy(
    "academic_year"
).count() \
.orderBy(
    "academic_year"
).show(
    truncate=False
)
 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

from pyspark.sql.functions import current_timestamp, lit
#df = spark.createDataFrame(df)
df = (
    df.withColumn("load_ts", current_timestamp())

)
 

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

spark.sql("""
CREATE SCHEMA IF NOT EXISTS lh_cs_silver.survey
""")

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

try:
    df.write \
    .format("delta") \
    .mode("overwrite") \
    .saveAsTable ("target_table")

    final_count = df.count()
    rows_read = final_count
    rows_written = final_count

except Exception as e:
    # updated audit log to FAILURE
    end_time = datetime.now()
    insert_pipeline_monitoring_record( 
        _safe_str(MASTER_RUN_ID), _safe_str(RUN_ID), _safe_str(PIPELINE_NAME), _safe_str(NOTEBOOK_NAME), _safe_str(DOMAIN), _safe_str(SOURCE_NAME), _safe_str(ACTIVITY_NAME),
        "failed", _safe_str(LEVEL), _safe_str(LAYER), _safe_str(source_table), _safe_str(target_table), 0, 0, 0,
        start_time, end_time, str(e), str(e), _safe_str(ENVIRONMENT), _safe_str(WORKSPACE_NAME), _safe_str(WORKSPACE_ID), _safe_str(PIPELINE_NAME),
        rerun_count=0
    )
    raise Exception(
        f"Failed to create/load survey summary data table: {e}"
    )



# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

end_time = datetime.now()
insert_pipeline_monitoring_record( 
    _safe_str(MASTER_RUN_ID), _safe_str(RUN_ID), _safe_str(PIPELINE_NAME), _safe_str(NOTEBOOK_NAME), _safe_str(DOMAIN), _safe_str(SOURCE_NAME), _safe_str(ACTIVITY_NAME),
    "success", _safe_str(LEVEL), _safe_str(LAYER), _safe_str(source_table), _safe_str(target_table), rows_read, rows_written, 0,
    start_time, end_time, "", "", _safe_str(ENVIRONMENT), _safe_str(WORKSPACE_NAME), _safe_str(WORKSPACE_ID), _safe_str(PIPELINE_NAME),
    rerun_count=0
)
print("SUCCESS logged")

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }
